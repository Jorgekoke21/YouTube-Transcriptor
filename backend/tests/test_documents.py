"""Renderers, filename sanitisation and cache fingerprint."""

from pathlib import Path

import pytest
from docx import Document

from app.models.document import DocumentType
from app.prompts import get_prompt_versions, prompt_version_string
from app.services.cache import compute_cache_key
from app.services.documents.builder import write_document_files
from app.services.documents.markdown import render_markdown
from app.services.documents.text import render_text
from app.utils.filenames import sanitize_filename
from tests.sample_document import sample_document

VID = "dQw4w9WgXcQ"


# ---------------------------------------------------------------- markdown


def test_markdown_structure_and_clickable_timestamps() -> None:
    md = render_markdown(sample_document())
    assert md.startswith("# Divergencias RSI\n")
    assert "> **Canal:** Canal de ejemplo" in md
    assert "> **Idioma de la transcripción:** Español · automática" in md
    assert "> **Tipo de documento:** Trading" in md
    assert "## Confirmación de entrada" in md
    assert "### Divergencia alcista" in md
    assert f"**Vídeo: [18:42 – 21:10 ↗](https://www.youtube.com/watch?v={VID}&t=1122s)**" in md
    assert "1. Detectar la divergencia" in md
    assert "- [ ] ¿Hay divergencia?" in md
    assert "> **Nota:** No se especifica una regla concreta" in md
    assert "| Tipo | Precio | RSI |" in md
    assert "Máximo \\| más bajo" in md  # pipes escaped in tables
    assert "## Ideas clave" in md


def test_markdown_english_labels() -> None:
    md = render_markdown(sample_document(output_language="en"))
    assert "**Video: [18:42 – 21:10 ↗]" in md
    assert "## Key points" in md and "> **Document type:** Trading" in md


def test_text_has_no_markdown_syntax() -> None:
    txt = render_text(sample_document())
    assert "**" not in txt and "](" not in txt
    assert "CONFIRMACIÓN DE ENTRADA" in txt
    assert "Vídeo: 18:42 – 21:10" in txt
    assert "[ ] ¿Hay divergencia?" in txt


def test_all_files_are_generated(tmp_path: Path) -> None:
    files = write_document_files(sample_document(), tmp_path / "out")
    for path in (files.processed_json, files.markdown, files.docx, files.pdf, files.txt):
        assert path.exists() and path.stat().st_size > 0
    assert files.pdf.read_bytes().startswith(b"%PDF")

    docx = Document(str(files.docx))
    headings = [(p.style.name, p.text) for p in docx.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert ("Title", "Divergencias RSI") in headings
    assert ("Heading 1", "Confirmación de entrada") in headings
    assert ("Heading 2", "Divergencia alcista") in headings
    assert len(docx.tables) == 2  # metadata + comparison table
    rels = [r.target_ref for r in docx.part.rels.values() if "hyperlink" in r.reltype]
    assert f"https://www.youtube.com/watch?v={VID}&t=1122s" in rels


# ---------------------------------------------------------------- filenames


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Divergencias RSI: cómo operarlas", "Divergencias_RSI_como_operarlas"),
        ("../../etc/passwd", "etc_passwd"),
        ("C:\\Windows\\system32", "C_Windows_system32"),
        ("CON", "document"),
        ("", "document"),
        (None, "document"),
        ("🔥🔥🔥", "document"),
        ("a" * 300, "a" * 80),
        ('Título "raro" <con> |pipes|?*', "Titulo_raro_con_pipes"),
    ],
)
def test_sanitize_filename(raw, expected) -> None:
    result = sanitize_filename(raw)
    assert result == expected
    assert "/" not in result and "\\" not in result and ".." not in result


# ---------------------------------------------------------------- cache


def test_cache_key_is_deterministic_and_sensitive_to_every_input() -> None:
    base = ("hash", "full_notes", "es", "base-v2+pipeline-v5+full-notes-v1", "openai:m")
    key = compute_cache_key(*base)
    assert key == compute_cache_key(*base) and len(key) == 64
    for i in range(len(base)):
        changed = list(base)
        changed[i] = changed[i] + "x"
        assert compute_cache_key(*changed) != key


def test_prompt_versions_are_tracked_per_mode() -> None:
    assert get_prompt_versions(DocumentType.TRADING) == {"base": "base-v2", "pipeline": "pipeline-v6", "mode": "trading-v3"}
    assert prompt_version_string(DocumentType.STUDY_GUIDE) == "base-v2+pipeline-v6+study-guide-v1"
    assert len({prompt_version_string(t) for t in DocumentType}) == len(DocumentType)
