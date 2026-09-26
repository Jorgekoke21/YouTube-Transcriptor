"""DOCX renderer (python-docx)."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from app.models.document import AnyDocument, ContentBlock, SourceRange
from app.services.documents.common import labels, metadata_lines, range_label, source_lines, split_bold, time_links

LINK_COLOR = "1D4ED8"
MUTED = RGBColor(0x6B, 0x72, 0x80)


def _add_hyperlink(paragraph: Paragraph, text: str, url: str, size: Pt | None = None) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    if size is not None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size.pt * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_rich_text(paragraph: Paragraph, text: str, italic: bool = False) -> None:
    for chunk, bold in split_bold(text):
        run = paragraph.add_run(chunk)
        run.bold = bold or None
        run.italic = italic or None


def _add_time_links(paragraph: Paragraph, doc: AnyDocument, ranges: Sequence[SourceRange], prefix: str = "") -> None:
    links = time_links(doc, ranges)
    if not links:
        return
    if prefix:
        run = paragraph.add_run(prefix)
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED
    for i, link in enumerate(links):
        if i:
            sep = paragraph.add_run(" · ")
            sep.font.size = Pt(9)
        _add_hyperlink(paragraph, link.label, link.url, size=Pt(9))


def _range_paragraph(doc_x, doc: AnyDocument, ranges: Sequence[SourceRange]) -> None:
    if not ranges:
        return
    p = doc_x.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    _add_time_links(p, doc, ranges, prefix=f"{range_label(doc)}: ")


def _list_paragraph(doc_x, marker: str, text: str) -> None:
    p = doc_x.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"{marker}\t")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
    _add_rich_text(p, text)


def _render_block(doc_x, block: ContentBlock, doc: AnyDocument) -> None:
    if block.type == "paragraph":
        p = doc_x.add_paragraph()
        _add_rich_text(p, block.text)
        if block.source_ranges:
            p.add_run(" ")
            _add_time_links(p, doc, block.source_ranges)
        return
    if block.type in ("quote", "note"):
        style = "Intense Quote" if block.type == "quote" else "Quote"
        p = doc_x.add_paragraph(style=style)
        if block.type == "note":
            r = p.add_run(f"{labels(doc.output_language)['note']}: ")
            r.bold = True
        _add_rich_text(p, block.text)
        if block.source_ranges:
            p.add_run(" ")
            _add_time_links(p, doc, block.source_ranges)
        return
    if block.type == "bullet_list":
        for item in block.items:
            _list_paragraph(doc_x, "•", item)
    elif block.type == "numbered_list":
        for n, item in enumerate(block.items, 1):
            _list_paragraph(doc_x, f"{n}.", item)
    elif block.type == "checklist":
        for item in block.items:
            _list_paragraph(doc_x, "☐", item)
    elif block.type == "table":
        width = max([len(block.table_headers)] + [len(r) for r in block.table_rows])
        rows = ([block.table_headers] if block.table_headers else []) + block.table_rows
        table = doc_x.add_table(rows=len(rows), cols=width)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_i, row in enumerate(rows):
            for c_i in range(width):
                cell = table.cell(r_i, c_i)
                cell.text = ""
                _add_rich_text(cell.paragraphs[0], row[c_i] if c_i < len(row) else "")
        doc_x.add_paragraph()
    if block.source_ranges:
        p = doc_x.add_paragraph()
        _add_time_links(p, doc, block.source_ranges)


def render_docx(doc: AnyDocument, path: Path) -> None:
    doc_x = Document()
    for page in doc_x.sections:
        page.left_margin = page.right_margin = Cm(2.3)
        page.top_margin = page.bottom_margin = Cm(2.2)
    normal = doc_x.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    lab = labels(doc.output_language)
    doc_x.core_properties.title = doc.title
    doc_x.add_heading(doc.title, level=0)

    meta = doc_x.add_table(rows=0, cols=2)
    meta.style = "Light List Accent 1"
    for m in metadata_lines(doc):
        cells = meta.add_row().cells
        label_run = cells[0].paragraphs[0].add_run(m.label)
        label_run.bold = True
        label_run.font.size = Pt(9)
        if m.url:
            _add_hyperlink(cells[1].paragraphs[0], m.value, m.url, size=Pt(9))
        else:
            run = cells[1].paragraphs[0].add_run(m.value)
            run.font.size = Pt(9)
    for row in meta.rows:
        row.cells[0].width = Cm(4.5)
        row.cells[1].width = Cm(11.5)
    doc_x.add_paragraph()

    sources = source_lines(doc)
    if sources:
        head = doc_x.add_paragraph()
        head.add_run(f"{lab['sources_used']}:").bold = True
        for src in sources:
            p = doc_x.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"{src.number}.\t")
            p.add_run(src.title).bold = True
            if src.channel:
                p.add_run(f" — {src.channel}")
            p.add_run(" — ")
            _add_hyperlink(p, src.url, src.url, size=Pt(9))
        doc_x.add_paragraph()

    if doc.summary.strip():
        doc_x.add_heading(lab["summary"], level=1)
        _add_rich_text(doc_x.add_paragraph(), doc.summary.strip())

    for section in doc.sections:
        if section.title.strip():
            doc_x.add_heading(section.title.strip(), level=1)
        _range_paragraph(doc_x, doc, section.source_ranges)
        for block in section.blocks:
            _render_block(doc_x, block, doc)
        for sub in section.subsections:
            doc_x.add_heading(sub.title.strip(), level=2)
            _range_paragraph(doc_x, doc, sub.source_ranges)
            for block in sub.blocks:
                _render_block(doc_x, block, doc)

    if doc.key_points:
        doc_x.add_heading(lab["key_points"], level=1)
        for kp in doc.key_points:
            p = doc_x.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.add_run("•\t")
            _add_rich_text(p, kp.text)
            if kp.source_ranges:
                p.add_run(" ")
                _add_time_links(p, doc, kp.source_ranges)

    last = doc_x.add_paragraph()
    last.add_run().add_break(WD_BREAK.LINE)
    doc_x.save(str(path))
